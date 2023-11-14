#! C:\Users\HP\AppData\Local\Programs\Python\Python311\python.exe

print("Content-Type: text/html\r\n\r\n")
import cgi
import mysql.connector
con=mysql.connector.connect(host='localhost', user='webrays', passwd='rayssoft',database='webrays')
t=con.cursor()
try:
    f=cgi.FieldStorage()
    d1=f.getvalue('t1')
    if d1=='search':
        t.execute('select * from rays_contact_us ORDER BY ai desc')
        rs=t.fetchall()
        # for i in rs:
        #     print(i[4])
        print('<table><thead><tr><th id="th1">Name</th><th id="th1">Subject</th><th id="th1">Mobile &nbsp; No.</th><th id="th1">Email &nbsp; No.</th><th id="th2">Message</th><th id="th1">Date & Time</th></tr></thead><tbody>')
        for i in rs:
            print('<tr><td>'+str(i[0])+'</td><td>'+str(i[1])+'</td><td>'+str(i[2])+'</td><td>'+str(i[3])+'</td><td><pre>'+str(i[4])+'</pre></td><td>'+str(i[6])+'</td></tr>')
        print('</tbody></table>')
        print('Hello')
    elif d1=='ent':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=f.getvalue('t7')
        d8=f.getvalue('t8')
        d9=f.getvalue('t9')
        url="insert into test (subject,topic,q,a,o1,o2,o3,o4) values(%s,%s,%s,%s,%s,%s,%s,%s)"
        t.execute(url,(d2,d3,d4,d5,d6,d7,d8,d9))
        con.commit()
        print('Successfully inserted!')
    elif d1=='testload':
        t.execute("select distinct subject from test")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='testload1':
        d2=f.getvalue('t2')
        t.execute("select distinct topic from test where subject='"+d2+"'")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='testsearch':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        t.execute("select * from test where subject='"+d2+"' and topic='"+d3+"'")
        rs=t.fetchall()
        if rs!=[]:
            print('<div style="text-align: right;"><br><br><input type="button" value="Update" id="b6"><input type="button" value="Delete" id="b7"></div><br><div style="height: 33rem; overflow-y: auto;"><table id="tab3"><tr id="sticky"><th>S. No.</th><th>Subject</th><th>Topic</th><th>Question</th><th>Answer</th><th>Option 1</th><th>Option 2</th><th>Option 3</th><th>Option 4</th><th>Check</th></tr>')
            for i in rs:
                # print('<tr><td>'+str(i[0])+'   '+'</td><td>'+str(i[1])+'   '+'</td><td>'+str(i[2])+'   '+'</td><td contenteditable="true" id="e">'+str(i[3])+'   '+'</td><td contenteditable="true" id="e">'+str(i[4])+'   '+'</td><td contenteditable="true" id="e">'+str(i[5])+'   '+'</td><td contenteditable="true" id="e">'+str(i[6])+'   '+'</td><td contenteditable="true" id="e">'+str(i[7])+'   '+'</td><td contenteditable="true" id="e">'+str(i[8])+'   '+'</td><td><input type="radio" id="r1" name="a"></td></tr>')
                print('<tr><td>'+str(i[8])+'   '+'</td><td>'+str(i[0])+'   '+'</td><td>'+str(i[1])+'   '+'</td><td contenteditable="true" id="e">'+str(i[2])+'   '+'</td><td contenteditable="true" id="e">'+str(i[3])+'   '+'</td><td contenteditable="true" id="e">'+str(i[4])+'   '+'</td><td contenteditable="true" id="e">'+str(i[5])+'   '+'</td><td contenteditable="true" id="e">'+str(i[6])+'   '+'</td><td contenteditable="true" id="e">'+str(i[7])+'   '+'</td><td><input type="radio" id="r1" name="a"></td></tr>')
            print('</table></div>')
        else:
            print(1)
    elif d1=='testup':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=f.getvalue('t7')
        d8=f.getvalue('t8')
        d9=f.getvalue('t9')
        d10=int(f.getvalue('t10'))
        url="update test set subject='%s',topic='%s',q='%s',a='%s',o1='%s',o2='%s',o3='%s',o4='%s' where sn=%d"%(d2,d3,d4,d5,d6,d7,d8,d9,d10)
        t.execute(url)
        con.commit()
        print('Updated Successfully!')
    elif d1=='testdl':
        d2=int(f.getvalue('t2'))
        url="delete from test where sn=%d"%(d2)
        t.execute(url)
        con.commit()
        print('Deleted Successfully!')
    elif d1=='exlup':
        import openpyxl
        # import pandas as pd
        # from xls2xlsx import XLS2XLSX as xx
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        # d5=xx(d4)
        # d5.to_xlsx('output.xlsx')
        # d5=pd.read_csv(d4,header=None,delim_whitespace=True)
        # d5.to_excel("output.xlsx",index=False,header=None)
        # print(d4)
        l=list()
        a=openpyxl.load_workbook(d4)
        b=a.active
        c=b.max_row
        url="insert into test (subject,topic,q,a,o1,o2,o3,o4) values(%s,%s,%s,%s,%s,%s,%s,%s)"
        for i in range(1,c+1):
            subject='C'
            topic='Function'
            q=b.cell(row=i,column=1).value
            a=b.cell(row=i,column=2).value
            o1=b.cell(row=i,column=3).value
            o2=b.cell(row=i,column=4).value
            o3=b.cell(row=i,column=5).value
            o4=b.cell(row=i,column=6).value
            val=(subject,topic,q,a,o1,o2,o3,o4)
            l.append(val)
        t.executemany(url,l)
        con.commit()
        print('Done!')
    elif d1=='load':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        a=set()
        t.execute("select * from test where subject='"+d2+"' and topic='"+d3+"' ORDER BY RAND() LIMIT 20")
        res=t.fetchall()
        # print(res)
        if len(res)!=20:
            print(0)
        else:
            for i in res:
                a.clear()
                a.add(i[4]),a.add(i[5]),a.add(i[6]),a.add(i[7])
                print(i[0],',,,,',i[1],',,,,',i[2],',,,,',i[3],',,,,')
                for j in a:
                    print(j)
                    print(',,,')
                print(',,,,,')
    elif d1=='ent1':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        url="insert into rays_proj_solution_video (subject,topic,technology,guide,link) values(%s,%s,%s,%s,%s)"
        t.execute(url,(d2,d3,d4,d5,d6))
        con.commit()
        print('Successfully inserted!')
    elif d1=='presload':
        t.execute("select distinct subject from rays_proj_solution_video")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='presload1':
        d2=f.getvalue('t2')
        t.execute("select distinct topic from rays_proj_solution_video where subject='"+d2+"'")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='pressearch':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        url="select * from rays_proj_solution_video where subject='"+d2+"'"
        if d3!=None:
            url=url+" and topic='"+d3+"'"
        t.execute(url)
        rs=t.fetchall()
        if rs!=[]:
            print('<div style="text-align: right;"><br><br><input type="button" value="Update" id="b6"><input type="button" value="Delete" id="b7"></div><br><div style="height: 33rem; overflow-y: auto;"><table id="tab3"><tr id="sticky"><th>S. No.</th><th>Subject</th><th>Topic</th><th>Guide</th><th>Technology</th><th>Link</th><th>Check</th></tr>')
            for i in rs:
                print('<tr><td>'+str(i[0])+'   '+'</td><td>'+str(i[1])+'   '+'</td><td contenteditable="true" id="e">'+str(i[2])+'   '+'</td><td contenteditable="true" id="e">'+str(i[3])+'   '+'</td><td contenteditable="true" id="e">'+str(i[4])+'   '+'</td><td contenteditable="true" id="e">'+str(i[5])+'   '+'</td><td><input type="radio" id="r1" name="a"></td></tr>')
            print('</table></div>')
        else:
            print(1)
    elif d1=='presup':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=int(f.getvalue('t7'))
        url="update rays_proj_solution_video set subject='%s',topic='%s',guide='%s',technology='%s',link='%s' where sn=%d"%(d2,d3,d4,d5,d6,d7)
        t.execute(url)
        con.commit()
        print('Updated Successfully!')
    elif d1=='presdl':
        d2=int(f.getvalue('t2'))
        url="delete from rays_proj_solution_video where sn=%d"%(d2)
        t.execute(url)
        con.commit()
        print('Deleted Successfully!')
    elif d1=='ent2':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=f.getvalue('t7')
        url="insert into rays_project_presentation (subject,topic,technology,tl,member,link) values(%s,%s,%s,%s,%s,%s)"
        t.execute(url,(d2,d3,d4,d5,d6,d7))
        con.commit()
        print('Successfully inserted!')
    elif d1=='pres1load':
        t.execute("select distinct subject from rays_project_presentation")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='pres1load1':
        d2=f.getvalue('t2')
        t.execute("select distinct topic from rays_project_presentation where subject='"+d2+"'")
        rs=t.fetchall()
        for i in rs:
            print('<option>'+i[0]+'</option>')
    elif d1=='pres1search':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        url="select * from rays_project_presentation where subject='"+d2+"'"
        if d3!=None:
            url=url+" and topic='"+d3+"'"
        t.execute(url)
        rs=t.fetchall()
        if rs!=[]:
            print('<div style="text-align: right;"><br><br><input type="button" value="Update" id="b6"><input type="button" value="Delete" id="b7"></div><br><div style="height: 33rem; overflow-y: auto;"><table id="tab3"><tr id="sticky"><th>S. No.</th><th>Subject</th><th>Topic</th><th>Technology</th><th>TL Name</th><th>Members</th><th>Link</th><th>Check</th></tr>')
            for i in rs:
                print('<tr><td>'+str(i[0])+'   '+'</td><td>'+str(i[1])+'   '+'</td><td contenteditable="true" id="e">'+str(i[2])+'   '+'</td><td contenteditable="true" id="e">'+str(i[3])+'   '+'</td><td contenteditable="true" id="e">'+str(i[4])+'   '+'</td><td contenteditable="true" id="e">'+str(i[5])+'   '+'</td><td contenteditable="true" id="e">'+str(i[6])+'   '+'</td><td><input type="radio" id="r1" name="a"></td></tr>')
            print('</table></div>')
        else:
            print(1)
    elif d1=='pres1up':
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=f.getvalue('t7')
        d8=int(f.getvalue('t8'))
        url="update rays_project_presentation set subject='%s',topic='%s',technology='%s',tl='%s',member='%s',link='%s' where sn=%d"%(d2,d3,d4,d5,d6,d7,d8)
        t.execute(url)
        con.commit()
        print('Updated Successfully!')
    elif d1=='pres1dl':
        d2=int(f.getvalue('t2'))
        url="delete from rays_project_presentation where sn=%d"%(d2)
        t.execute(url)
        con.commit()
        print('Deleted Successfully!')
    elif d1=='Development':
        d2=f.getvalue('t2')
        t.execute("select * from rays_project_presentation where subject='"+d2+"'")
        rs1=t.fetchall()
        t.execute("select * from rays_proj_solution_video where subject='"+d2+"'")
        rs2=t.fetchall()
        for i in rs1:
            print('<div id="x1">'+i[6]+'<h3 align="center">'+i[4]+'</h3><p id="p1"><b>Member : </b><label>'+i[5]+'</label></p><hr><p id="p1"><b>Topic : </b><label>'+i[2]+'</label></p><hr><p id="p1"><b>Used Technology : </b><label>'+i[3]+'</label></p></div>')
        print(',,,')
        for i in rs2:
            print('<div id="x2">'+i[5]+'<h3 align="center"><b>Guide By : </b><label>'+i[2]+'</label></h3><p id="p1"><b>Topic : </b><label>'+i[3]+'</label></p><p id="p1"><b>Technology : </b><label>'+i[4]+'</label></p></div>')
    elif d1=='carear':
        from datetime import date
        a=date.today()
        t.execute('select * from rays_carear order by sn desc')
        rs=t.fetchall()
        for i in rs:
            x=i[5].split('-')
            b=date(int(x[0]),int(x[1]),int(x[2]))
            c=(b-a).days
            if c>=0:
                print('<div id="parent"><div id="d1"><p><b>Position : <label>'+i[0]+'</label></b></p><p><i class="fas fa-clock"></i> &nbsp; <label>'+i[2]+'</label></p><p><i class="fas fa-book-open"></i> &nbsp; <label>Skills Required</label></p><p style="color: blue;"> &nbsp; &nbsp; &nbsp;  &nbsp; <label>'+i[1]+'</label></p><p><i class="far fa-clock"></i> &nbsp; Closing Date : <label>'+x[2]+'-'+x[1]+'-'+x[0]+' ('+str(c)+' Days Left)</label></p></div>')
                print('<div id="d2"><p style="text-align: center;"><input type="button" value="Apply Now" id="b1"></p><p><i class="fas fa-users" id="iright"></i> &nbsp; <label>'+i[4]+'</label> vacancies</p><p><i class="fas fa-cubes"></i> &nbsp; <label>'+i[3]+'</label> Experience</p><p><i class="fa fa-map-marker"></i> &nbsp; &nbsp; Place : <label>'+i[6]+'</label></p><p><i class="fas fa-donate"></i>&nbsp;  &nbsp; Salary : <label>'+i[7]+'</label></p><p style="text-align: center;"><i class="fas fa-check" id="status" style="color: blue;"></i> &nbsp; <label style="color: #07ac93 ;">Active</label></p></div></div>')
        for i in rs:
            x=i[5].split('-')
            b=date(int(x[0]),int(x[1]),int(x[2]))
            c=(b-a).days
            if c<0:
                print('<div id="parent"><div id="d1"><p><b>Position : <label>'+i[0]+'</label></b></p><p><i class="fas fa-clock"></i> &nbsp; <label>'+i[2]+'</label></p><p><i class="fas fa-book-open"></i> &nbsp; <label>Skills Required</label></p><p style="color: blue;"> &nbsp; &nbsp; &nbsp;  &nbsp; <label>'+i[1]+'</label></p><p><i class="far fa-clock"></i> &nbsp; Closing Date : <label>'+x[2]+'-'+x[1]+'-'+x[0]+' (Expire)</label></p></div>')
                print('<div id="d2"><p style="text-align: center;"><br><p><i class="fas fa-users" id="iright"></i> &nbsp; <label>'+i[4]+'</label> vacancies</p><p><i class="fas fa-cubes"></i> &nbsp; <label>'+i[3]+'</label> Experience</p><p><i class="fa fa-map-marker"></i> &nbsp;  &nbsp; Place : <label>'+i[6]+'</label></p><p><i class="fas fa-donate"></i>&nbsp;  &nbsp; Salary : <label>'+i[7]+'</label></p><p style="text-align: center; color:red;">X &nbsp; <label>Inactive</label></p></div></div>')
    elif d1=='job_apply':
        d2=f.getvalue('t2').title()
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        d7=f.getvalue('t7')
        d8=f.getvalue('t8')
        if d8=='entern':
            url="insert into rays_intern_apply (name,mo_no,email,resume,position,mode) values(%s,%s,%s,%s,%s,%s)"
        else:
            url="insert into rays_job_apply (name,mo_no,email,resume,position,mode) values(%s,%s,%s,%s,%s,%s)"
        t.execute(url,(d2,d3,d4,d5,d6,d7))
        con.commit()
        print('Thanks for apply!\nOur team will Contact You soon')
    elif d1=='job_ser':
        t.execute('select * from rays_job_apply order by sn desc')
        rs=t.fetchall()
        print('<table id="t1"><tr><th>Name</th><th>Mo_No</th><th>Email</th><th>Position</th><th>Mode</th><th id="th">Resume</th></tr>')
        for i in rs:
            a=str(i[4]).split("'")
            print('<tr><td>'+i[1]+'</td><td>'+i[2]+'</td><td>'+i[3]+'</td><td>'+i[5]+'</td><td>'+i[6]+'</td><td><embed src="'+a[1]+'" type="application/pdf"/></td></tr>')
        print('</table>')
    elif d1=='intern':
        dd=f.getvalue('tt')
        if dd=='fl':
            d2=f.getvalue('t2')
            d3=f.getvalue('t3')
            d4=f.getvalue('t4')
            d5=f.getvalue('t5')
            d6=f.getvalue('t6')
            d7=int(f.getvalue('t7'))
            url='select * from internship_details where stipend>=%d'%(d7)
            if d2!=None:
                url=url+" and position='"+d2+"'"
            if d3!=None:
                url=url+" and location='"+d3+"'"
            if d4!=None:
                url1="'"+d4+"'"
            if d5!=None and d4!=None:
                url1=url1+",'"+d5+"'"
            if d5!=None and d4==None:
                url1="'"+d5+"'"
            if d6!=None and (d5!=None or d4!=None):
                url1=url1+",'"+d6+"'"
            if d6!=None and d5==None and d4==None:
                url1="'"+d6+"'"
            if(d4!=None or d5!=None or d6!=None):
                url=url+" and category in ("+url1+")"
            url=url+" order by sn desc"
            t.execute(url)
        else:
            t.execute('select * from internship_details order by sn desc')
        rs=t.fetchall()
        from datetime import date
        a=date.today()
        if(rs!=[]):
            for i in rs:
                x=i[8].split('-')
                y=i[4].split('-')
                b=date(int(x[0]),int(x[1]),int(x[2]))
                c=(b-a).days
                if c>=0:
                    print('<div id="d1"><div id="x"><img src="REPL_t-shirt_logo.png" alt="loading..."></div><p style="text-align:center;"><b>'+i[1]+'</b></p><p><i class="fas fa-map-marker-alt"></i>&nbsp; &nbsp;<label> '+i[3]+'</label></p><p><i class="far fa-play-circle"></i>&nbsp; &nbsp;<label>Start &nbsp; Date &nbsp; <b>:</b> &nbsp; </label><label>'+y[2]+'-'+y[1]+'-'+y[0]+'</label></p><p><i class="fas fa-calendar-alt"></i>&nbsp; &nbsp;<label>Duration &nbsp; <b>:</b> &nbsp; </label><label>'+i[5]+'</label></p><p><i class="fas fa-donate"></i>&nbsp; &nbsp;<label>Stipend &nbsp; <b>:</b> &nbsp; </label>&#8377; '+str(i[6])+'/<label id="m">'+i[2]+'</label></p>')
                    print('<p><i class="fas fa-book-reader"></i>&nbsp; &nbsp;<label>Certificate Available &nbsp; <b>:</b> &nbsp; </label><label>Yes</label></p><p><i class="fas fa-user-graduate"></i>&nbsp; &nbsp;<label>Skill Required &nbsp; <b>:</b> &nbsp; </label><label>'+i[7]+'</label></p><p><i class="fas fa-bell-slash"></i>&nbsp;&nbsp;<label>Last Apply Date <b>:</b> </label>'+x[2]+'-'+x[1]+'-'+x[0]+' ('+str(c)+' Days Left)</p></br><p id="p1"><input type="button" value="Apply Now" id="b2"></p></div>')
            for i in rs:
                x=i[8].split('-')
                b=date(int(x[0]),int(x[1]),int(x[2]))
                c=(b-a).days
                if c<0:
                    print('<div id="d1"><div id="x"><img src="REPL_t-shirt_logo.png" alt="loading..."></div><p style="text-align:center;"><b>'+i[1]+'</b></p><p><i class="fas fa-map-marker-alt"></i>&nbsp; &nbsp;<label> '+i[3]+'</label></p><p><i class="far fa-play-circle"></i>&nbsp; &nbsp;<label>Start &nbsp; Date &nbsp; <b>:</b> &nbsp; </label><label>'+y[2]+'-'+y[1]+'-'+y[0]+'</label></p><p><i class="fas fa-calendar-alt"></i>&nbsp; &nbsp;<label>Duration &nbsp; <b>:</b> &nbsp; </label><label>'+i[5]+'</label></p><p><i class="fas fa-donate"></i>&nbsp; &nbsp;<label>Stipend &nbsp; <b>:</b> &nbsp; </label>&#8377; '+str(i[6])+'/<label id="m">'+i[2]+'</label></p>')
                    print('<p><i class="fas fa-book-reader"></i>&nbsp; &nbsp;<label>Certificate Available &nbsp; <b>:</b> &nbsp; </label><label>Yes</label></p><p><i class="fas fa-user-graduate"></i>&nbsp; &nbsp;<label>Skill Required &nbsp; <b>:</b> &nbsp; </label><label>'+i[7]+'</label></p><p><i class="fas fa-bell-slash"></i>&nbsp;&nbsp;<label>Last Apply Date <b>:</b> </label>'+x[2]+'-'+x[1]+'-'+x[0]+' (<label style="color:red;">Expired</label>)</p></div>')
        else:
            print(5)
        if dd=='lo':
            a1=set();a2=set()
            for i in rs:
                a1.add(i[1])
                a2.add(i[3])
            print(',,,,,')
            for i in a1:
                print('<option>'+i+'</option>')
            print(',,,,,')
            for i in a2:
                print('<option>'+i+'</option>')
    else:
        d2=f.getvalue('t2')
        d3=f.getvalue('t3')
        d4=f.getvalue('t4')
        d5=f.getvalue('t5')
        d6=f.getvalue('t6')
        # print(d5)
        url='insert into rays_contact_us (name,subject,mo_no,email,message,d_t) values(%s,%s,%s,%s,%s,%s)'
        t.execute(url,(d1,d2,d3,d4,d5,d6))
        con.commit()
        print('Successes')
except Exception as e:
    print('Unsuccesses',e)